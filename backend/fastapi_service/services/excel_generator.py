# excel_generator.py — Excel Output Generator (v4.0 — FIXED)
# ─────────────────────────────────────────────────────────────────────────────
# Fixes v4.0:
#   1. Title row shows the real source filename (strips job-id UUID prefix that
#      was leaking into the title when file_name was the stored upload path).
#   2. _flatten_record: handles list-wrapped values, None, floats, and empty
#      strings uniformly — no crash on any input type.
#   3. Numeric columns cast safely; NaN shown as blank (not "nan").
#   4. _create_summary_sheet: confidence safely averaged even for mixed types.
#   5. Auto-creates output directory.
#   6. Works for direct PDF records, OCR records, and ITC-register records.
# ─────────────────────────────────────────────────────────────────────────────

import os
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("gst2_fastapi.excel_generator")

# ── Column definitions ────────────────────────────────────────────────────────
COLUMNS = [
    ("GSTIN of Supplier",         "GSTIN",              False),
    ("Trade / Legal Name",        "TRADE_NAME",         False),
    ("Invoice Number",            "INVOICE_NO",         False),
    ("Invoice Type",              "INVOICE_TYPE",       False),
    ("Invoice Date",              "INVOICE_DATE",       False),
    ("Invoice Value (₹)",         "INVOICE_VALUE",      True),
    ("Place of Supply",           "PLACE_OF_SUPPLY",    False),
    ("Reverse Charge",            "REVERSE_CHARGE",     False),
    ("Tax Rate (%)",              "TAX_RATE",           True),
    ("Taxable Value (₹)",         "TAXABLE_VALUE",      True),
    ("IGST (₹)",                  "IGST",               True),
    ("CGST (₹)",                  "CGST",               True),
    ("SGST / UTGST (₹)",          "SGST",               True),
    ("CESS (₹)",                  "CESS",               True),
    ("GSTR-1 Filing Status",      "GSTR1_STATUS",       False),
    ("GSTR-1 Filing Date",        "GSTR1_FILING_DATE",  False),
    ("GSTR-1 Period",             "GSTR1_PERIOD",       False),
    ("GSTR-3B Filing Status",     "GSTR3B_STATUS",      False),
    ("Amendment",                 "AMENDMENT",          False),
    ("Tax Period Amended",        "TAX_PERIOD_AMENDED", False),
    ("Cancellation Date",         "CANCEL_DATE",        False),
    ("Source",                    "SOURCE",             False),
    ("IRN",                       "IRN",                False),
    ("IRN Date",                  "IRN_DATE",           False),
]

NUMERIC_KEYS = {col[1] for col in COLUMNS if col[2]}

# ── Styling palette ───────────────────────────────────────────────────────────
COLOR_HEADER_BG = "1F4E79"
COLOR_HEADER_FG = "FFFFFF"
COLOR_TITLE_BG  = "2E75B6"
COLOR_ALT_ROW   = "EBF3FB"
COLOR_TOTAL_BG  = "1F4E79"
COLOR_SUB_HDR   = "D6E4F0"
COLOR_BORDER    = "B8CCE4"

# UUID pattern to strip from display names
_UUID_PREFIX = re.compile(
    r'^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}_',
    re.IGNORECASE
)


def _display_name(file_name: str) -> str:
    """Strip job-ID UUID prefix from filename for cleaner Excel titles."""
    stem = Path(file_name).stem
    # Remove leading UUID like "9982a939-eca1-48dd-bce6-010bf8dee05d_"
    stem = _UUID_PREFIX.sub('', stem)
    return stem


# ─────────────────────────────────────────────────────────────────────────────
# PUBLIC API
# ─────────────────────────────────────────────────────────────────────────────

class ExcelGenerator:

    def generate(
        self,
        extracted_data: List[Dict[str, Any]],
        output_path: str,
        file_name: str,
        job_id: str,
    ) -> str:
        output_dir = Path(output_path)
        output_dir.mkdir(parents=True, exist_ok=True)

        timestamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
        disp_name  = _display_name(file_name)
        xlsx_name  = f"GST2A_{disp_name}_{timestamp}_{job_id[:8]}.xlsx"
        xlsx_path  = output_dir / xlsx_name

        logger.info(f"Generating Excel: {xlsx_name}  ({len(extracted_data)} records)")

        rows        = [self._flatten_record(r) for r in extracted_data]
        col_keys    = [c[1] for c in COLUMNS]
        col_headers = [c[0] for c in COLUMNS]

        df = pd.DataFrame(rows, columns=col_keys) if rows else pd.DataFrame(columns=col_keys)
        df.columns = col_headers

        for col_def in COLUMNS:
            disp, key, is_num = col_def
            if is_num and disp in df.columns:
                df[disp] = pd.to_numeric(df[disp], errors="coerce")

        df.to_excel(str(xlsx_path), index=False, sheet_name="GST 2A Data", engine="openpyxl")
        self._apply_style(str(xlsx_path), file_name, len(rows))
        self._create_summary_sheet(str(xlsx_path), extracted_data, file_name)

        logger.info(f"Excel saved: {xlsx_path}")
        return str(xlsx_path)

    # ─────────────────────────────────────────────────────────────────────────
    # INTERNAL HELPERS
    # ─────────────────────────────────────────────────────────────────────────

    def _flatten_record(self, rec: Dict[str, Any]) -> Dict[str, Any]:
        """
        Normalise a record dict:
        - Unwrap single-element lists
        - Convert None / NaN to ""
        - Keep numeric columns as numeric (or "" for blanks)
        """
        flat = {}
        for _, key, is_num in COLUMNS:
            raw = rec.get(key, "")
            # Unwrap list
            if isinstance(raw, list):
                raw = raw[0] if raw else ""
            # Handle None / nan
            if raw is None:
                raw = ""
            if is_num:
                # Keep numeric as-is (pd.to_numeric will handle "" → NaN)
                flat[key] = raw
            else:
                s = str(raw).strip()
                flat[key] = "" if s.lower() in ("none", "nan", "") else s
        return flat

    def _apply_style(self, xlsx_path: str, file_name: str, row_count: int):
        wb = load_workbook(xlsx_path)
        ws = wb.active
        ws.title = "GST 2A Data"

        thin   = Side(style="thin",   color=COLOR_BORDER)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.insert_rows(1, amount=2)
        last_col = get_column_letter(len(COLUMNS))

        # Row 1: Title — FIX: use clean display name, not raw filename with UUID
        title_cell       = ws["A1"]
        title_cell.value = f"GST FORM 2A — {_display_name(file_name)}"
        title_cell.font  = Font(name="Calibri", bold=True, size=14, color=COLOR_HEADER_FG)
        title_cell.fill  = PatternFill("solid", fgColor=COLOR_TITLE_BG)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"A1:{last_col}1")
        ws.row_dimensions[1].height = 28

        # Row 2: Subtitle
        sub_cell       = ws["A2"]
        sub_cell.value = (
            f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}    "
            f"│   Records: {row_count}    │   Source: GSTN Portal"
        )
        sub_cell.font  = Font(name="Calibri", italic=True, size=9, color="444444")
        sub_cell.fill  = PatternFill("solid", fgColor=COLOR_SUB_HDR)
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"A2:{last_col}2")
        ws.row_dimensions[2].height = 16

        # Row 3: Column headers
        for cell in ws[3]:
            cell.font      = Font(name="Calibri", bold=True, size=9, color=COLOR_HEADER_FG)
            cell.fill      = PatternFill("solid", fgColor=COLOR_HEADER_BG)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = border
        ws.row_dimensions[3].height = 38

        # Data rows (alternate shading)
        for row_idx, row in enumerate(ws.iter_rows(min_row=4), 1):
            for cell in row:
                cell.font      = Font(name="Calibri", size=8)
                cell.border    = border
                cell.alignment = Alignment(vertical="center")
                if row_idx % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=COLOR_ALT_ROW)

        # Column widths
        widths = [
            22, 22, 16, 10, 13, 14, 18, 13, 9,
            14, 12, 12, 14, 10,
            15, 16, 13, 15,
            11, 15, 16, 14, 66, 13
        ]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Totals row
        data_start = 4
        data_end   = ws.max_row
        total_row  = data_end + 2

        tc = ws.cell(total_row, 1, "TOTALS")
        tc.font = Font(name="Calibri", bold=True, color=COLOR_HEADER_FG, size=9)
        tc.fill = PatternFill("solid", fgColor=COLOR_TOTAL_BG)
        tc.alignment = Alignment(horizontal="center")

        for col_num in [6, 10, 11, 12, 13, 14]:
            letter = get_column_letter(col_num)
            cell   = ws.cell(total_row, col_num)
            cell.value          = f"=SUM({letter}{data_start}:{letter}{data_end})"
            cell.number_format  = "#,##0.00"
            cell.font           = Font(name="Calibri", bold=True, color=COLOR_HEADER_FG, size=9)
            cell.fill           = PatternFill("solid", fgColor=COLOR_TOTAL_BG)
            cell.alignment      = Alignment(horizontal="right")

        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{last_col}3"

        wb.save(xlsx_path)

    def _create_summary_sheet(
        self,
        xlsx_path: str,
        extracted_data: List[Dict[str, Any]],
        file_name: str,
    ):
        try:
            wb = load_workbook(xlsx_path)

            if "Summary" in wb.sheetnames:
                del wb["Summary"]

            ws = wb.create_sheet("Summary")

            thin   = Side(style="thin", color=COLOR_BORDER)
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            rows = [self._flatten_record(r) for r in extracted_data]
            df   = pd.DataFrame(rows, columns=[c[1] for c in COLUMNS]) if rows else pd.DataFrame(columns=[c[1] for c in COLUMNS])

            for _, key, is_num in COLUMNS:
                if is_num:
                    df[key] = pd.to_numeric(df[key], errors="coerce")

            # Safe confidence average (handles float, str, list, None)
            raw_conf = [r.get("_confidence", 0) for r in extracted_data]
            conf_vals = []
            for v in raw_conf:
                if isinstance(v, list):
                    v = v[0] if v else 0
                try:
                    conf_vals.append(float(v))
                except (TypeError, ValueError):
                    conf_vals.append(0.0)
            avg_conf = (sum(conf_vals) / len(conf_vals) * 100) if conf_vals else 0.0

            total_records    = len(extracted_data)
            total_inv_value  = df["INVOICE_VALUE"].sum()
            total_taxable    = df["TAXABLE_VALUE"].sum()
            total_igst       = df["IGST"].sum()
            total_cgst       = df["CGST"].sum()
            total_sgst       = df["SGST"].sum()
            total_cess       = df["CESS"].sum()
            total_tax        = total_igst + total_cgst + total_sgst + total_cess
            unique_suppliers = df["GSTIN"].nunique()

            def fmt_amt(v):
                if pd.isna(v) or v == 0:
                    return "₹0.00"
                return f"₹{v:,.2f}"

            source_counts = df["SOURCE"].value_counts().to_dict() if "SOURCE" in df.columns else {}
            source_str    = ", ".join(f"{k}:{v}" for k, v in source_counts.items()) or "—"

            summary_data = [
                ["EXTRACTION SUMMARY", "", ""],
                ["", "", ""],
                ["Metric", "Value", "Notes"],
                ["Total Invoice Records",   total_records,              "Rows extracted"],
                ["Unique Suppliers",        unique_suppliers,           "Distinct GSTINs"],
                ["Total Invoice Value",     fmt_amt(total_inv_value),   "Sum of all invoices"],
                ["Total Taxable Value",     fmt_amt(total_taxable),     "Taxable amount"],
                ["Total IGST",             fmt_amt(total_igst),        "Integrated GST"],
                ["Total CGST",             fmt_amt(total_cgst),        "Central GST"],
                ["Total SGST/UTGST",       fmt_amt(total_sgst),        "State/UT GST"],
                ["Total CESS",             fmt_amt(total_cess),        "Cess amount"],
                ["Total Tax Collected",    fmt_amt(total_tax),         "IGST+CGST+SGST+CESS"],
                ["", "", ""],
                ["PROCESSING INFO", "", ""],
                ["", "", ""],
                ["Source File",            _display_name(file_name),   "Uploaded document"],
                ["Source Types",           source_str,                 "E-Invoice/GSTR-1/Manual..."],
                ["OCR Accuracy",           f"{avg_conf:.1f}%",         "Average confidence"],
                ["Generated On",           datetime.now().strftime("%d/%m/%Y %H:%M:%S"), ""],
            ]

            for r_idx, row_data in enumerate(summary_data, 1):
                for c_idx, val in enumerate(row_data, 1):
                    cell       = ws.cell(r_idx, c_idx, val)
                    cell.font  = Font(name="Calibri", size=10)
                    cell.border = border

            # Title styling
            title = ws.cell(1, 1)
            title.font  = Font(name="Calibri", bold=True, size=13, color=COLOR_HEADER_FG)
            title.fill  = PatternFill("solid", fgColor=COLOR_TITLE_BG)
            title.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells("A1:C1")
            ws.row_dimensions[1].height = 26

            # Header row (row 3)
            for cell in ws[3]:
                cell.font  = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
                cell.fill  = PatternFill("solid", fgColor=COLOR_HEADER_BG)
                cell.alignment = Alignment(horizontal="center")

            # Section header (PROCESSING INFO)
            for r_idx, row_data in enumerate(summary_data, 1):
                if row_data[0] == "PROCESSING INFO":
                    cell = ws.cell(r_idx, 1)
                    cell.font  = Font(name="Calibri", bold=True, size=10, color=COLOR_HEADER_FG)
                    cell.fill  = PatternFill("solid", fgColor=COLOR_TITLE_BG)
                    ws.merge_cells(f"A{r_idx}:C{r_idx}")

            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 26
            ws.column_dimensions["C"].width = 30

            wb.save(xlsx_path)
            logger.info("Summary sheet added successfully")

        except Exception as e:
            logger.error(f"Summary sheet creation failed: {e}", exc_info=True)


# ── Singleton ─────────────────────────────────────────────────────────────────
_excel_gen_instance: Optional[ExcelGenerator] = None


def get_excel_generator() -> ExcelGenerator:
    global _excel_gen_instance
    if _excel_gen_instance is None:
        _excel_gen_instance = ExcelGenerator()
    return _excel_gen_instance